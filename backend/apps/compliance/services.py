"""Compliance domain: contracts, SSF filing, compliance status, contractor summaries."""

from __future__ import annotations

import csv
import datetime as dt
import io

from django.conf import settings
from django.core.files.base import ContentFile
from django.db import transaction
from django.db.models import Count, Q
from django.utils import timezone
from django.utils.translation import get_language
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from apps.core.bs import MONTHS_EN, MONTHS_NE, BSMonth, bs_str, js_weekday, to_devanagari, today_ad
from apps.core.exceptions import APIError, InvalidState
from apps.core.money import EXCEL_LAKH_FORMAT, D, money_str, q2

from .models import EmploymentContract, SSFFiling

WEEKDAYS_EN = ["Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]
WEEKDAYS_NE = ["आइतबार", "सोमबार", "मङ्गलबार", "बुधबार", "बिहीबार", "शुक्रबार", "शनिबार"]

# Changing any of these creates a new contract version.
VERSIONED_TERMS = (
    "designation",
    "site_id",
    "wage_type",
    "basic_salary",
    "dearness_allowance",
    "other_allowance",
    "daily_wage",
)


# ---------------------------------------------------------------------------
# Contracts
# ---------------------------------------------------------------------------


def _m(value):
    return money_str(value) if value is not None else None


def contract_terms(emp) -> dict:
    org, worker = emp.organization, emp.worker
    shift = emp.shift
    return {
        "org_name": org.name,
        "org_name_ne": org.name_ne,
        "org_address": org.address,
        "org_district": org.district.name_en if org.district_id else None,
        "org_pan": org.pan_no,
        "worker_name": worker.full_name,
        "worker_name_ne": worker.full_name_ne,
        "citizenship_no": worker.citizenship_no,
        "designation": emp.designation,
        "trade": emp.trade.name_en if emp.trade_id else None,
        "site_id": str(emp.site_id) if emp.site_id else None,
        "site_name": emp.site.name if emp.site_id else None,
        "start_date": emp.start_date.isoformat(),
        "start_date_bs": bs_str(emp.start_date),
        "wage_type": emp.wage_type,
        "basic_salary": _m(emp.basic_salary),
        "dearness_allowance": _m(emp.dearness_allowance),
        "other_allowance": _m(emp.other_allowance or 0),
        "daily_wage": _m(emp.daily_wage),
        "monthly_gross": _m(emp.monthly_gross),
        "ssf_enrolled": emp.ssf_enrolled,
        "shift": {
            "name": shift.name,
            "start_time": shift.start_time.strftime("%H:%M"),
            "end_time": shift.end_time.strftime("%H:%M"),
        }
        if shift
        else None,
        "weekly_off_en": WEEKDAYS_EN[org.weekly_off_day],
        "weekly_off_ne": WEEKDAYS_NE[org.weekly_off_day],
    }


def latest_contract(emp) -> EmploymentContract | None:
    return emp.contracts.order_by("-version").first()


def terms_changed(contract: EmploymentContract, terms: dict) -> bool:
    return any(contract.terms.get(k) != terms.get(k) for k in VERSIONED_TERMS)


@transaction.atomic
def generate_contract(emp, user=None, language="both", force=False) -> tuple[EmploymentContract, bool]:
    """Create a new version when terms changed (or forced); otherwise return the latest."""
    terms = contract_terms(emp)
    latest = latest_contract(emp)
    if latest and not force and not terms_changed(latest, terms) and latest.language == language:
        return latest, False
    version = (latest.version + 1) if latest else 1
    contract = EmploymentContract.objects.create(
        employment=emp, version=version, generated_by=user, language=language, terms=terms
    )
    return contract, True


def accept_contract(contract: EmploymentContract, channel="app") -> EmploymentContract:
    if contract.accepted_at is None:
        contract.accepted_at = timezone.now()
        contract.accepted_channel = channel
        contract.save(update_fields=["accepted_at", "accepted_channel"])
    return contract


# ---------------------------------------------------------------------------
# SSF filing
# ---------------------------------------------------------------------------

SSF_COLUMNS = [
    "ssf_id",
    "worker_name",
    "citizenship_no",
    "basic_earned",
    "employee_11",
    "employer_20",
    "total_31",
    "bs_year",
    "bs_month",
]
MISSING_COLUMNS = ["worker_name", "phone", "citizenship_no", "designation", "basic_earned", "total_31"]


def ssf_due_date(run) -> dt.date:
    """BS month end + SSF_DUE_DAYS days."""
    return run.period_end + dt.timedelta(days=settings.SSF_DUE_DAYS)


def ssf_rows(run) -> tuple[list[dict], list[dict]]:
    """(rows with an SSF ID, rows missing an SSF ID) for SSF-enrolled payslips."""
    rows, missing = [], []
    slips = run.payslips.select_related("worker", "employment").filter(employment__ssf_enrolled=True)
    for s in slips.order_by("worker__full_name"):
        total = q2(s.ssf_employee + s.ssf_employer)
        base = {
            "worker_name": s.worker.full_name,
            "citizenship_no": s.worker.citizenship_no or "",
            "basic_earned": money_str(s.basic_earned),
            "total_31": money_str(total),
        }
        if s.worker.ssf_id:
            rows.append(
                {
                    "ssf_id": s.worker.ssf_id,
                    **base,
                    "employee_11": money_str(s.ssf_employee),
                    "employer_20": money_str(s.ssf_employer),
                    "bs_year": run.bs_year,
                    "bs_month": run.bs_month,
                }
            )
        else:
            missing.append(
                {
                    **base,
                    "phone": s.worker.phone,
                    "designation": s.employment.designation,
                    "employment_id": str(s.employment_id),
                    "worker_id": str(s.worker_id),
                }
            )
    return rows, missing


def ssf_summary(rows, missing) -> dict:
    def total(key, src):
        return money_str(sum((D(r[key]) for r in src), start=D(0)))

    return {
        "rows": len(rows),
        "missing": len(missing),
        "basic": total("basic_earned", rows),
        "employee_11": total("employee_11", rows),
        "employer_20": total("employer_20", rows),
        "total_31": total("total_31", rows),
    }


def ssf_xlsx(run) -> bytes:
    rows, missing = ssf_rows(run)
    wb = Workbook()
    ws = wb.active
    ws.title = "SSF contributions"
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="EEF3FF")
    org = run.organization
    ws.append(
        [
            f"{org.name} — SSF contributions {MONTHS_EN[run.bs_month]} {run.bs_year} "
            f"({MONTHS_NE[run.bs_month]} {run.bs_year}) · SSF employer no. {org.ssf_employer_no or '—'}"
        ]
    )
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(SSF_COLUMNS)
    for cell in ws[3]:
        cell.font, cell.fill = bold, fill
    for r in rows:
        ws.append(
            [
                r["ssf_id"],
                r["worker_name"],
                r["citizenship_no"],
                float(r["basic_earned"]),
                float(r["employee_11"]),
                float(r["employer_20"]),
                float(r["total_31"]),
                r["bs_year"],
                r["bs_month"],
            ]
        )
        for col in (4, 5, 6, 7):
            ws.cell(row=ws.max_row, column=col).number_format = EXCEL_LAKH_FORMAT
    for i, w in enumerate([16, 26, 18, 14, 14, 14, 14, 8, 8], start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    ms = wb.create_sheet("Missing SSF ID")
    ms.append(["Workers enrolled in SSF without an SSF ID — register them on the SSF portal first."])
    ms["A1"].font = bold
    ms.append([])
    ms.append(MISSING_COLUMNS)
    for cell in ms[3]:
        cell.font, cell.fill = bold, fill
    for m in missing:
        ms.append(
            [
                m["worker_name"],
                m["phone"],
                m["citizenship_no"],
                m["designation"],
                float(m["basic_earned"]),
                float(m["total_31"]),
            ]
        )
    for i, w in enumerate([26, 14, 18, 22, 14, 14], start=1):
        ms.column_dimensions[chr(64 + i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ssf_filing_csv(run) -> str:
    rows, _missing = ssf_rows(run)
    out = io.StringIO()
    w = csv.DictWriter(out, fieldnames=SSF_COLUMNS, extrasaction="ignore")
    w.writeheader()
    w.writerows(rows)
    return out.getvalue()


def ssf_filename(run, ext) -> str:
    return f"ssf-filing-{run.bs_year}-{run.bs_month:02d}-{MONTHS_EN[run.bs_month].lower()}.{ext}"


@transaction.atomic
def prepare_ssf_filing(run, user=None) -> tuple[SSFFiling, bool]:
    if run.status not in ("approved", "paid"):
        raise InvalidState(message=_("SSF filing can be prepared only for approved or paid payroll runs."))
    filing = SSFFiling.objects.select_for_update().filter(run=run).first()
    if filing and filing.status != "prepared":
        raise InvalidState(message=_("This SSF filing is already submitted."))
    rows, missing = ssf_rows(run)
    content = ContentFile(ssf_xlsx(run), name=ssf_filename(run, "xlsx"))
    created = filing is None
    if created:
        filing = SSFFiling(run=run)
    elif filing.file:
        filing.file.delete(save=False)
    filing.due_date = ssf_due_date(run)
    filing.prepared_at = timezone.now()
    filing.prepared_by = user
    filing.summary = ssf_summary(rows, missing)
    filing.file = content
    filing.save()
    return filing, created


FILING_ORDER = {"prepared": 0, "submitted": 1, "acknowledged": 2}


def update_filing(filing: SSFFiling, data: dict) -> SSFFiling:
    status = data.get("status", filing.status)
    if FILING_ORDER[status] < FILING_ORDER[filing.status]:
        raise InvalidState(message=_("An SSF filing cannot move back to an earlier status."))
    if status in ("submitted", "acknowledged") and not (data.get("submitted_on") or filing.submitted_on):
        data["submitted_on"] = today_ad()
    if status == "acknowledged" and not (data.get("receipt_no") or filing.receipt_no):
        raise APIError("validation_error", fields={"receipt_no": [_("Enter the SSF receipt number.")]})
    filing.status = status
    for f in ("submitted_on", "receipt_no", "note"):
        if f in data:
            setattr(filing, f, data[f])
    filing.save()
    return filing


def runs_needing_ssf_filing(org=None):
    """Approved/paid runs with SSF contributions whose filing is not yet submitted."""
    from apps.payroll.models import PayrollRun

    qs = PayrollRun.objects.filter(status__in=["approved", "paid"], payslips__ssf_employee__gt=0)
    if org is not None:
        qs = qs.filter(organization=org)
    qs = qs.filter(Q(ssf_filing__isnull=True) | Q(ssf_filing__status="prepared"))
    return qs.distinct().select_related("organization")


# ---------------------------------------------------------------------------
# Compliance status
# ---------------------------------------------------------------------------

WEIGHTS = {
    "ssf_employer_registered": 15,
    "workers_missing_ssf_id": 10,
    "ssf_filings_overdue": 15,
    "below_minimum_wage": 15,
    "contracts_missing": 10,
    "contracts_unaccepted": 5,
    "attendance_unmarked_7d": 10,
    "payroll_not_run_last_month": 15,
    "credentials_expiring": 5,
}


def _item(key, status, count, message, link):
    return {"key": key, "status": status, "count": count, "message": str(message), "link": link}


def _below_minimum(emps, on) -> int:
    from apps.hiring.models import MinimumWage

    mw = MinimumWage.current(on)
    if not mw:
        return 0
    n = 0
    for e in emps:
        if e.wage_type == "monthly" and D(e.basic_salary) + D(e.dearness_allowance) < mw.monthly_total:
            n += 1
        elif e.wage_type == "daily" and D(e.daily_wage) < mw.daily:
            n += 1
    return n


def _unmarked_days(org, emps, today) -> tuple[int, int]:
    """(missing worker-days, expected worker-days) over the 7 days before today."""
    from apps.workforce.models import Attendance, Holiday

    start, end = today - dt.timedelta(days=7), today - dt.timedelta(days=1)
    holidays = set(
        Holiday.objects.filter(
            Q(organization__isnull=True) | Q(organization=org), date__range=(start, end)
        ).values_list("date", flat=True)
    )
    marked = set(
        Attendance.objects.filter(employment__in=emps, date__range=(start, end)).values_list(
            "employment_id", "date"
        )
    )
    expected = missing = 0
    day = start
    while day <= end:
        if day not in holidays and js_weekday(day) != org.weekly_off_day:
            for e in emps:
                if e.is_active_on(day):
                    expected += 1
                    if (e.id, day) not in marked:
                        missing += 1
        day += dt.timedelta(days=1)
    return missing, expected


def expiring_credentials_count(org, days=30) -> int:
    try:
        from apps.credentials.models import Credential
    except Exception:  # noqa: BLE001 - credentials app not built yet
        return 0
    today = today_ad()
    try:
        return (
            Credential.objects.filter(
                worker__employments__organization=org,
                worker__employments__status="active",
                expires_on__isnull=False,
                expires_on__range=(today, today + dt.timedelta(days=days)),
            )
            .exclude(status__in=["rejected", "expired"])
            .distinct()
            .count()
        )
    except Exception:  # noqa: BLE001 - schema not migrated yet
        return 0


def compliance_status(org) -> dict:
    from apps.payroll.models import PayrollRun
    from apps.workforce.models import Employment

    today = today_ad()
    active = list(Employment.objects.filter(organization=org, status="active").select_related("worker"))
    active_ids = [e.id for e in active]
    items = []

    # SSF employer registration
    ok = bool(org.ssf_employer_no)
    items.append(
        _item(
            "ssf_employer_registered",
            "ok" if ok else "action",
            0 if ok else 1,
            _("Your organization is registered with SSF.")
            if ok
            else _("Add your SSF employer registration number in organization settings."),
            "/app/settings",
        )
    )

    # Workers without SSF ID
    n = sum(1 for e in active if e.ssf_enrolled and not e.worker.ssf_id)
    items.append(
        _item(
            "workers_missing_ssf_id",
            "warning" if n else "ok",
            n,
            _("%(count)s SSF-enrolled workers have no SSF ID.") % {"count": n}
            if n
            else _("Every SSF-enrolled worker has an SSF ID."),
            "/app/workers",
        )
    )

    # SSF filings overdue
    pending = list(runs_needing_ssf_filing(org))
    overdue = [r for r in pending if ssf_due_date(r) < today]
    due_soon = [r for r in pending if today <= ssf_due_date(r) <= today + dt.timedelta(days=5)]
    if overdue:
        status, n, msg = (
            "action",
            len(overdue),
            _("%(count)s SSF filings are overdue.") % {"count": len(overdue)},
        )
    elif due_soon:
        status, n = "warning", len(due_soon)
        msg = _("%(count)s SSF filings are due within 5 days.") % {"count": n}
    else:
        status, n, msg = "ok", 0, _("SSF filings are up to date.")
    items.append(_item("ssf_filings_overdue", status, n, msg, "/app/compliance#ssf"))

    # Minimum wage
    n = _below_minimum(active, today)
    items.append(
        _item(
            "below_minimum_wage",
            "action" if n else "ok",
            n,
            _("%(count)s workers are paid below the minimum wage.") % {"count": n}
            if n
            else _("Everyone is paid at least the minimum wage."),
            "/app/workers",
        )
    )

    # Contracts
    with_contract = set(
        EmploymentContract.objects.filter(employment_id__in=active_ids).values_list(
            "employment_id", flat=True
        )
    )
    n = len([i for i in active_ids if i not in with_contract])
    items.append(
        _item(
            "contracts_missing",
            "action" if n else "ok",
            n,
            _("%(count)s workers have no written contract.") % {"count": n}
            if n
            else _("Every worker has a written contract."),
            "/app/workers",
        )
    )
    latest = {}
    for c in EmploymentContract.objects.filter(employment_id__in=active_ids).order_by(
        "employment_id", "version"
    ):
        latest[c.employment_id] = c
    n = sum(1 for c in latest.values() if c.accepted_at is None)
    items.append(
        _item(
            "contracts_unaccepted",
            "warning" if n else "ok",
            n,
            _("%(count)s contracts are waiting for the worker's acceptance.") % {"count": n}
            if n
            else _("All contracts are accepted."),
            "/app/workers",
        )
    )

    # Hajiri in the last 7 days
    missing, expected = _unmarked_days(org, active, today)
    status = "ok" if not missing else ("action" if expected and missing / expected > 0.2 else "warning")
    items.append(
        _item(
            "attendance_unmarked_7d",
            status,
            missing,
            _("%(count)s worker-days of hajiri are not marked in the last 7 days.") % {"count": missing}
            if missing
            else _("Hajiri is marked for the last 7 days."),
            "/app/attendance",
        )
    )

    # Payroll last month
    prev = BSMonth.of(today).previous()
    had_staff = (
        Employment.objects.filter(organization=org, start_date__lte=prev.end)
        .filter(Q(end_date__isnull=True) | Q(end_date__gte=prev.start))
        .exists()
    )
    run = PayrollRun.objects.filter(organization=org, bs_year=prev.year, bs_month=prev.month).first()
    month = (
        f"{MONTHS_NE[prev.month]} {to_devanagari(prev.year)}"
        if get_language() == "ne"
        else f"{MONTHS_EN[prev.month]} {prev.year}"
    )
    if not had_staff or (run and run.status != "draft"):
        status, n, msg = "ok", 0, _("Payroll for %(month)s is done.") % {"month": month}
    elif run:
        status, n, msg = "warning", 1, _("Payroll for %(month)s is still a draft.") % {"month": month}
    else:
        status, n, msg = "action", 1, _("Payroll for %(month)s has not been run.") % {"month": month}
    items.append(_item("payroll_not_run_last_month", status, n, msg, "/app/payroll"))

    # Credentials expiring
    n = expiring_credentials_count(org)
    items.append(
        _item(
            "credentials_expiring",
            "warning" if n else "ok",
            n,
            _("%(count)s worker certificates expire within 30 days.") % {"count": n}
            if n
            else _("No certificates expire in the next 30 days."),
            "/app/compliance#credentials",
        )
    )

    penalty = sum(
        WEIGHTS[i["key"]] * (1 if i["status"] == "action" else 0.5 if i["status"] == "warning" else 0)
        for i in items
    )
    return {"score": max(0, round(100 - penalty)), "items": items}


# ---------------------------------------------------------------------------
# Contractors
# ---------------------------------------------------------------------------


def contractor_summary(contractor, month: BSMonth) -> dict:
    from apps.payroll.models import Payslip
    from apps.workforce.models import Attendance, Employment

    start, end = month.start, month.end
    emps = list(
        Employment.objects.filter(contractor_assignment__contractor=contractor, start_date__lte=end)
        .filter(Q(end_date__isnull=True) | Q(end_date__gte=start))
        .select_related("worker")
        .order_by("worker__full_name")
    )
    slips = {
        s.employment_id: s
        for s in Payslip.objects.filter(
            run__organization=contractor.organization,
            run__bs_year=month.year,
            run__bs_month=month.month,
            employment__in=emps,
        ).select_related("run")
    }
    att = {}
    for row in (
        Attendance.objects.filter(
            employment__in=emps, date__range=(start, end), status__in=["present", "half_day"]
        )
        .values("employment_id", "status")
        .annotate(n=Count("id"))
    ):
        weight = D("1") if row["status"] == "present" else D("0.5")
        att[row["employment_id"]] = att.get(row["employment_id"], D(0)) + weight * row["n"]
    workers, days = [], D(0)
    gross = net = employer = D(0)
    for e in emps:
        slip = slips.get(e.id)
        d = D(slip.present_days) if slip else att.get(e.id, D(0))
        days += d
        if slip:
            gross += slip.gross_pay
            net += slip.net_pay
            employer += slip.gross_pay + slip.ssf_employer
        workers.append(
            {
                "employment_id": str(e.id),
                "worker_name": e.worker.full_name,
                "designation": e.designation,
                "days": str(d),
                "gross_pay": money_str(slip.gross_pay) if slip else None,
                "net_pay": money_str(slip.net_pay) if slip else None,
            }
        )
    run = next(iter(slips.values())).run if slips else None
    return {
        "contractor": {"id": str(contractor.id), "name": contractor.name},
        "bs_year": month.year,
        "bs_month": month.month,
        "month_name_en": MONTHS_EN[month.month],
        "month_name_ne": MONTHS_NE[month.month],
        "headcount": len(emps),
        "days": str(days),
        "wages": {
            "gross": money_str(gross),
            "net": money_str(net),
            "employer_cost": money_str(employer),
        }
        if slips
        else None,
        "payroll_run": {"id": str(run.id), "status": run.status} if run else None,
        "workers": workers,
    }
