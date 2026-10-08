"""Payout provider adapters (docs/CONTRACT-ECOSYSTEM.md §6).

``BankFileProvider`` is the working, always-available fallback: it produces a bulk-transfer file
that the employer uploads to their own bank, then reconciles the bank's result file.
Wallet / ConnectIPS adapters stay disabled until a merchant agreement exists: they raise
``NotConfigured`` unless their code is in ``PAYOUT_PROVIDERS`` *and* credentials are set.
"""

from __future__ import annotations

import csv
import hashlib
import hmac
import io
import json
from dataclasses import dataclass

from django.conf import settings
from django.core.files.base import ContentFile
from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Font

from apps.core.bs import MONTHS_EN
from apps.core.exceptions import APIError
from apps.core.money import EXCEL_LAKH_FORMAT, money_str


class NotConfigured(APIError):
    default_code = "provider_not_configured"

    def __init__(self, message=None):
        super().__init__(message=message or _("This payout provider is not configured."))


class InvalidSignature(APIError):
    status_code = 403
    default_code = "invalid_signature"

    def __init__(self, message=None):
        super().__init__(message=message or "Invalid webhook signature.")


class WebhookReplayed(APIError):
    status_code = 409
    default_code = "webhook_replayed"

    def __init__(self):
        super().__init__(message="This webhook delivery was already processed.")


@dataclass
class ItemUpdate:
    status: str  # sent | paid | failed
    item_id: str | None = None
    provider_ref: str | None = None
    failure_reason: str | None = None


BANK_FILE_COLUMNS = ["name", "account", "bank", "amount", "narration", "reference"]
PAID_WORDS = {"paid", "success", "successful", "completed", "complete", "ok", "done", "credited"}
FAILED_WORDS = {"failed", "failure", "rejected", "returned", "error", "invalid", "reversed"}


def narration(batch) -> str:
    run = batch.run
    return f"Salary {MONTHS_EN[run.bs_month]} {run.bs_year}"


class Provider:
    code = ""
    methods: frozenset = frozenset()
    credential_settings: tuple = ()

    def is_enabled(self) -> bool:
        # ``PAYOUT_PROVIDERS`` is the default; a ``payout_<code>`` feature flag overrides it.
        from apps.core.flags import is_enabled

        return is_enabled(f"payout_{self.code}") and self.has_credentials()

    def has_credentials(self) -> bool:
        return all(getattr(settings, name, None) for name in self.credential_settings)

    def ensure_enabled(self):
        if not self.is_enabled():
            raise NotConfigured(
                _("%(provider)s payouts are not enabled yet. Use the bank file instead.")
                % {"provider": self.label}
            )

    @property
    def label(self) -> str:
        from .models import PROVIDER_LABELS

        return str(PROVIDER_LABELS.get(self.code, self.code))

    def submit(self, batch) -> None:
        raise NotImplementedError

    # -------------------------------------------------------------- webhooks
    def webhook_secret(self) -> str | None:
        return (settings.PAYOUT_WEBHOOK_SECRETS or {}).get(self.code)

    def verify_signature(self, request) -> None:
        """HMAC-SHA256 in ``X-Signature`` (hex, optional ``sha256=`` prefix).

        With ``X-Timestamp`` (unix seconds) the signed message is ``ts + "." + body`` and the
        timestamp must be within ``PAYOUT_WEBHOOK_MAX_SKEW_SECONDS``; without it the raw body is
        signed (legacy providers). A verified signature is remembered for
        ``PAYOUT_WEBHOOK_REPLAY_TTL_SECONDS``: the exact same delivery again → 409 ``webhook_replayed``.
        """
        secret = self.webhook_secret()
        if not secret:
            raise NotConfigured()
        ts = (request.headers.get("X-Timestamp") or "").strip()
        message = request.body
        if ts:
            try:
                ts_value = int(ts)
            except ValueError as exc:
                raise InvalidSignature("Invalid X-Timestamp.") from exc
            skew = abs(timezone.now().timestamp() - ts_value)
            if skew > int(getattr(settings, "PAYOUT_WEBHOOK_MAX_SKEW_SECONDS", 300)):
                raise InvalidSignature("Webhook timestamp is outside the allowed window.")
            message = ts.encode() + b"." + request.body
        expected = hmac.new(secret.encode(), message, hashlib.sha256).hexdigest()
        given = (request.headers.get("X-Signature") or "").strip()
        if given.lower().startswith("sha256="):
            given = given[7:]
        if not given or not hmac.compare_digest(expected, given.lower()):
            raise InvalidSignature()
        request._ks_webhook_replay_key = self.remember_delivery(given.lower())

    def remember_delivery(self, signature: str) -> str:
        """Reject an exact replay of a verified delivery (cache ``add`` is atomic). Returns the key
        so the view can forget it when processing fails (the provider's retry must still work)."""
        from django.core.cache import cache

        digest = hashlib.sha256(f"{self.code}:{signature}".encode()).hexdigest()
        key = f"ks:payout-webhook:{digest}"
        ttl = int(getattr(settings, "PAYOUT_WEBHOOK_REPLAY_TTL_SECONDS", 86400))
        if not cache.add(key, 1, ttl):
            raise WebhookReplayed()
        return key

    def parse_webhook(self, request) -> list[ItemUpdate]:
        """Verify the HMAC and parse ``{"items": [{reference, status, provider_ref?, reason?}]}``."""
        self.verify_signature(request)
        try:
            payload = json.loads(request.body or b"{}")
        except ValueError as exc:
            raise APIError("validation_error", message="Invalid JSON body.") from exc
        raw_items = payload.get("items") if isinstance(payload, dict) else None
        if raw_items is None and isinstance(payload, dict):
            raw_items = [payload]
        updates = []
        for raw in raw_items or []:
            if not isinstance(raw, dict):
                continue
            status = classify_status(raw.get("status"))
            if status is None:
                continue
            updates.append(
                ItemUpdate(
                    status=status,
                    item_id=raw.get("reference") or raw.get("item_id"),
                    provider_ref=raw.get("provider_ref") or raw.get("transaction_id"),
                    failure_reason=raw.get("reason") or raw.get("failure_reason"),
                )
            )
        return updates


def classify_status(value) -> str | None:
    word = str(value or "").strip().lower()
    if word in PAID_WORDS:
        return "paid"
    if word in FAILED_WORDS:
        return "failed"
    if word in {"sent", "pending", "processing", "queued"}:
        return "sent"
    return None


class BankFileProvider(Provider):
    code = "bank_file"
    methods = frozenset({"bank"})

    def is_enabled(self) -> bool:
        return self.code in settings.PAYOUT_PROVIDERS

    def rows(self, batch):
        text = narration(batch)
        for item in batch.items.select_related("payslip__worker").order_by("payslip__worker__full_name"):
            yield [
                item.account_name or item.payslip.worker.full_name,
                item.account_no or "",
                item.bank_name or "",
                money_str(item.amount),
                text,
                str(item.id),
            ]

    def csv_bytes(self, batch) -> bytes:
        out = io.StringIO()
        w = csv.writer(out)
        w.writerow(BANK_FILE_COLUMNS)
        w.writerows(self.rows(batch))
        return out.getvalue().encode("utf-8")

    def xlsx_bytes(self, batch) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Bulk transfer"
        ws.append(BANK_FILE_COLUMNS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in self.rows(batch):
            row[3] = float(row[3])
            ws.append(row)
            ws.cell(row=ws.max_row, column=4).number_format = EXCEL_LAKH_FORMAT
        for col, width in zip("ABCDEF", [28, 22, 26, 14, 22, 38], strict=True):
            ws.column_dimensions[col].width = width
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def submit(self, batch) -> None:
        run = batch.run
        name = f"bank-transfer-{run.bs_year}-{run.bs_month:02d}-{str(batch.id)[:8]}.csv"
        if batch.file:
            batch.file.delete(save=False)
        batch.file = ContentFile(self.csv_bytes(batch), name=name)
        batch.items.filter(status="pending").update(status="sent", updated_at=timezone.now())

    def parse_webhook(self, request):
        raise NotConfigured(_("Bank file payouts have no webhook; upload the bank's result file instead."))

    def parse_result_file(self, upload) -> list[dict]:
        """Bank result CSV → [{row, reference, account, amount, status, provider_ref, reason}]."""
        raw = upload.read()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise APIError("validation_error", fields={"file": [_("The file is empty.")]})
        headers = {h: (h or "").strip().lower().replace(" ", "_") for h in reader.fieldnames}

        def pick(row, *names):
            for original, norm in headers.items():
                if norm in names and row.get(original) not in (None, ""):
                    return str(row[original]).strip()
            return None

        out = []
        for i, row in enumerate(reader, start=2):
            out.append(
                {
                    "row": i,
                    "reference": pick(row, "reference", "ref", "item_id", "customer_reference"),
                    "account": pick(row, "account", "account_no", "account_number"),
                    "amount": pick(row, "amount", "amount_npr"),
                    "status": classify_status(pick(row, "status", "result", "transaction_status")),
                    "provider_ref": pick(row, "bank_ref", "provider_ref", "transaction_id", "txn_id", "utr"),
                    "reason": pick(row, "reason", "remarks", "failure_reason", "message"),
                }
            )
        return out


class _WalletStub(Provider):
    """Live API adapter placeholder — disabled until a merchant agreement exists."""

    def submit(self, batch) -> None:
        self.ensure_enabled()
        # Live API call goes here once the merchant agreement and sandbox credentials exist.
        # Until then an enabled provider only marks items as sent and waits for the webhook.
        batch.items.filter(status="pending").update(status="sent", updated_at=timezone.now())


class EsewaProvider(_WalletStub):
    code = "esewa"
    methods = frozenset({"esewa"})
    credential_settings = ("ESEWA_MERCHANT_CODE", "ESEWA_SECRET_KEY")


class KhaltiProvider(_WalletStub):
    code = "khalti"
    methods = frozenset({"khalti"})
    credential_settings = ("KHALTI_SECRET_KEY",)


class ConnectIPSProvider(_WalletStub):
    code = "connectips"
    methods = frozenset({"bank"})
    credential_settings = ("CONNECTIPS_MERCHANT_ID", "CONNECTIPS_APP_ID", "CONNECTIPS_CREDITOR_KEY_PATH")


PROVIDERS: dict[str, Provider] = {
    p.code: p for p in (BankFileProvider(), EsewaProvider(), KhaltiProvider(), ConnectIPSProvider())
}


def get_provider(code: str) -> Provider:
    provider = PROVIDERS.get(code)
    if provider is None:
        raise APIError("validation_error", fields={"provider": [f"Unknown provider '{code}'."]})
    return provider
