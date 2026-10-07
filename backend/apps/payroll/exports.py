"""Payroll exports: wage sheet (xlsx), bank transfer (csv), SSF (csv), payslip (pdf)."""

from __future__ import annotations

import csv
import io

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.core.bs import MONTHS_EN, MONTHS_NE, bs_str, format_bs_en
from apps.core.money import EXCEL_LAKH_FORMAT, D, format_npr, lakh_group, q2

WAGE_SHEET_COLUMNS = [
    # (header_en, header_ne, getter, kind)
    ("S.N.", "क्र.सं.", None, "int"),
    ("Worker name", "कामदारको नाम", lambda p: p.worker.full_name, "text"),
    ("Designation", "पद", lambda p: p.employment.designation, "text"),
    ("Wage type", "तलब प्रकार", lambda p: p.wage_type, "text"),
    ("Join date (BS)", "सुरु मिति (वि.सं.)", lambda p: bs_str(p.employment.start_date), "text"),
    ("Days in month", "महिनाका दिन", lambda p: p.days_in_month, "int"),
    ("Present days", "उपस्थित दिन", lambda p: p.present_days, "num1"),
    ("Paid leave", "सशुल्क बिदा", lambda p: p.paid_leave_days, "num1"),
    ("Holidays", "सार्वजनिक बिदा", lambda p: p.holiday_days, "num1"),
    ("Weekly off", "साप्ताहिक बिदा", lambda p: p.weekly_off_days, "num1"),
    ("Absent", "अनुपस्थित", lambda p: p.absent_days, "num1"),
    ("Payable days", "भुक्तानी दिन", lambda p: p.payable_days, "num1"),
    ("OT hours", "ओभरटाइम घण्टा", lambda p: p.overtime_hours, "num2"),
    ("Basic salary", "आधारभूत तलब", lambda p: p.basic_earned, "money"),
    ("Dearness allowance", "महँगी भत्ता", lambda p: p.da_earned, "money"),
    ("Other allowance", "अन्य भत्ता", lambda p: p.allowance_earned, "money"),
    ("Overtime pay", "ओभरटाइम रकम", lambda p: p.overtime_pay, "money"),
    ("Dashain allowance", "चाडपर्व खर्च", lambda p: p.festival_allowance, "money"),
    ("Gross pay", "कुल तलब", lambda p: p.gross_pay, "money"),
    ("SSF employee 11%", "सा.सु.कोष कर्मचारी ११%", lambda p: p.ssf_employee, "money"),
    ("SSF employer 20%", "सा.सु.कोष रोजगारदाता २०%", lambda p: p.ssf_employer, "money"),
    ("TDS", "अग्रिम कर कट्टी", lambda p: p.tds, "money"),
    ("Advance recovery", "पेस्की कट्टी", lambda p: p.advance_recovery, "money"),
    ("Other deductions", "अन्य कट्टी", lambda p: p.other_deductions, "money"),
    ("Net pay", "खुद भुक्तानी", lambda p: p.net_pay, "money"),
    ("Payment method", "भुक्तानी माध्यम", lambda p: p.payment_method, "text"),
    ("Signature", "हस्ताक्षर", lambda p: "", "text"),
]


def _payslips(run):
    return list(run.payslips.select_related("worker", "employment").order_by("worker__full_name"))


def wage_sheet_xlsx(run) -> bytes:
    org = run.organization
    slips = _payslips(run)
    wb = Workbook()
    ws = wb.active
    ws.title = f"{MONTHS_EN[run.bs_month]} {run.bs_year}"
    ncols = len(WAGE_SHEET_COLUMNS)
    last_col = get_column_letter(ncols)
    bold = Font(bold=True)
    thin = Side(style="thin", color="C4CAD4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="EEF3FF")

    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = f"{org.name}" + (f" / {org.name_ne}" if org.name_ne else "")
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = (
        f"Wage sheet / तलब पाना — {MONTHS_EN[run.bs_month]} {run.bs_year} "
        f"({MONTHS_NE[run.bs_month]} {run.bs_year})"
        f"   Period (BS): {bs_str(run.period_start)} to {bs_str(run.period_end)}"
        f"   (AD {run.period_start.isoformat()} – {run.period_end.isoformat()})"
    )
    ws.merge_cells(f"A3:{last_col}3")
    ws["A3"] = (
        f"PAN: {org.pan_no or '—'}   SSF employer no.: {org.ssf_employer_no or '—'}   "
        f"Status: {run.status.upper()}   Generated: {format_bs_en(timezone.localdate())}"
    )
    header_row = 5
    for col, (en, ne, _, _) in enumerate(WAGE_SHEET_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=f"{en}\n{ne}")
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[header_row].height = 34

    row = header_row
    for idx, slip in enumerate(slips, start=1):
        row += 1
        for col, (_, _, getter, kind) in enumerate(WAGE_SHEET_COLUMNS, start=1):
            value = idx if getter is None else getter(slip)
            if kind in ("money", "num1", "num2"):
                value = float(value)
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = border
            if kind == "money":
                cell.number_format = EXCEL_LAKH_FORMAT
            elif kind == "num1":
                cell.number_format = "0.0"
            elif kind == "num2":
                cell.number_format = "0.00"

    total_row = row + 1
    ws.cell(row=total_row, column=2, value="Total / जम्मा").font = bold
    for col, (_, _, _, kind) in enumerate(WAGE_SHEET_COLUMNS, start=1):
        if kind == "money":
            # Static (pre-computed) totals: correct in any viewer, no formula evaluation needed.
            total = sum((D(WAGE_SHEET_COLUMNS[col - 1][2](s)) for s in slips), start=D(0))
            cell = ws.cell(row=total_row, column=col, value=float(q2(total)))
            cell.number_format = EXCEL_LAKH_FORMAT
            cell.font = bold
            cell.border = border
    totals_text_row = total_row + 2
    net_total = sum((D(s.net_pay) for s in slips), start=D(0))
    ws.cell(
        row=totals_text_row, column=2, value=f"Net payable (जम्मा खुद भुक्तानी): {format_npr(net_total)}"
    ).font = bold
    ws.cell(row=totals_text_row + 2, column=2, value="Prepared by / तयार गर्ने: ____________")
    ws.cell(row=totals_text_row + 2, column=10, value="Approved by / स्वीकृत गर्ने: ____________")

    widths = [6, 24, 16, 10, 14] + [10] * 8 + [13] * 12 + [12, 14]
    for i, w in enumerate(widths[:ncols], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def bank_transfer_csv(run) -> str:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["name", "account", "bank", "amount"])
    for s in _payslips(run):
        if s.payment_method != "bank":
            continue
        w.writerow(
            [
                s.worker.bank_account_name or s.worker.full_name,
                s.worker.bank_account_no or "",
                s.worker.bank_name or "",
                f"{q2(s.net_pay):.2f}",
            ]
        )
    return out.getvalue()


def ssf_csv(run) -> str:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["ssf_id", "name", "basic", "employee_11", "employer_20", "total_31"])
    for s in _payslips(run):
        if not s.employment.ssf_enrolled:
            continue
        w.writerow(
            [
                s.worker.ssf_id or "",
                s.worker.full_name,
                f"{q2(s.basic_earned):.2f}",
                f"{q2(s.ssf_employee):.2f}",
                f"{q2(s.ssf_employer):.2f}",
                f"{q2(s.ssf_employee + s.ssf_employer):.2f}",
            ]
        )
    return out.getvalue()


def payslip_pdf(slip) -> bytes:
    run = slip.run
    org = run.organization
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"Payslip {slip.worker.full_name} {MONTHS_EN[run.bs_month]} {run.bs_year}",
    )
    styles = getSampleStyleSheet()
    brand = colors.HexColor("#2647C4")
    story = [
        Paragraph(f"<font color='#2647C4'><b>{org.name}</b></font>", styles["Title"]),
        Paragraph(
            f"{org.address or ''} {org.district.name_en if org.district_id else ''}"
            f"{' · PAN ' + org.pan_no if org.pan_no else ''}",
            styles["Normal"],
        ),
        Spacer(1, 4 * mm),
        Paragraph(f"<b>Payslip — {MONTHS_EN[run.bs_month]} {run.bs_year} B.S.</b>", styles["Heading2"]),
        Paragraph(
            f"Period: {bs_str(run.period_start)} to {bs_str(run.period_end)} B.S. "
            f"({run.period_start.isoformat()} to {run.period_end.isoformat()} A.D.)",
            styles["Normal"],
        ),
        Spacer(1, 4 * mm),
    ]
    emp = slip.employment
    info = [
        ["Worker", slip.worker.full_name, "Designation", emp.designation or "—"],
        ["Employee code", emp.employee_code or "—", "Joined", f"{bs_str(emp.start_date)} B.S."],
        ["Wage type", slip.wage_type.title(), "SSF ID", slip.worker.ssf_id or "—"],
        ["Days in month", str(slip.days_in_month), "Payable days", str(slip.payable_days)],
        [
            "Present",
            str(slip.present_days),
            "Paid leave / holiday / weekly off",
            f"{slip.paid_leave_days} / {slip.holiday_days} / {slip.weekly_off_days}",
        ],
        ["Absent", str(slip.absent_days), "Overtime hours", str(slip.overtime_hours)],
    ]
    t = Table(info, colWidths=[32 * mm, 52 * mm, 50 * mm, 40 * mm])
    t.setStyle(
        TableStyle(
            [
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#6B7383")),
                ("TEXTCOLOR", (2, 0), (2, -1), colors.HexColor("#6B7383")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story += [t, Spacer(1, 5 * mm)]

    earnings = [
        ("Basic salary", slip.basic_earned),
        ("Dearness allowance", slip.da_earned),
        ("Other allowance", slip.allowance_earned),
        ("Overtime", slip.overtime_pay),
        ("Dashain allowance", slip.festival_allowance),
    ]
    deductions = [
        ("SSF (11%)", slip.ssf_employee),
        ("TDS", slip.tds),
        ("Advance (peski)", slip.advance_recovery),
        ("Other deductions", slip.other_deductions),
    ]
    rows = [["Earnings", "Rs.", "Deductions", "Rs."]]
    for i in range(max(len(earnings), len(deductions))):
        e = earnings[i] if i < len(earnings) else ("", None)
        d = deductions[i] if i < len(deductions) else ("", None)
        rows.append(
            [
                e[0],
                lakh_group(e[1]) if e[1] is not None else "",
                d[0],
                lakh_group(d[1]) if d[1] is not None else "",
            ]
        )
    total_ded = slip.ssf_employee + slip.tds + slip.advance_recovery + slip.other_deductions
    rows.append(["Gross pay", lakh_group(slip.gross_pay), "Total deductions", lakh_group(total_ded)])
    t2 = Table(rows, colWidths=[52 * mm, 35 * mm, 52 * mm, 35 * mm])
    t2.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), brand),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ("ALIGN", (3, 0), (3, -1), "RIGHT"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDE1E8")),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
            ]
        )
    )
    story += [t2, Spacer(1, 5 * mm)]
    story.append(Paragraph(f"<b>Net pay: {format_npr(slip.net_pay)}</b>", styles["Heading2"]))
    story.append(
        Paragraph(
            f"Employer SSF contribution (20%, not deducted): {format_npr(slip.ssf_employer)}",
            styles["Normal"],
        )
    )
    status = "PAID" if slip.payment_status == "paid" else "PENDING"
    paid = f" on {bs_str(timezone.localtime(slip.paid_at).date())} B.S." if slip.paid_at else ""
    story.append(Paragraph(f"Payment: {slip.payment_method.title()} — {status}{paid}", styles["Normal"]))
    story += [
        Spacer(1, 10 * mm),
        Paragraph(
            "<font size=8 color='#6B7383'>Generated by KaamSathi. "
            "This is a computer-generated payslip.</font>",
            styles["Normal"],
        ),
    ]
    doc.build(story)
    return buf.getvalue()
