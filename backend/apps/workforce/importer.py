"""Employment import from CSV/XLSX and the downloadable template."""

from __future__ import annotations

import csv
import datetime as dt
import io
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from apps.core.bs import BSDateError, parse_date_flexible
from apps.core.exceptions import APIError
from apps.core.phone import is_valid_phone, normalize_phone

from .models import Employment
from .services import create_employment, get_or_create_worker, validate_wage

COLUMNS = [
    "full_name",
    "phone",
    "designation",
    "start_date",
    "date_system",
    "wage_type",
    "basic_salary",
    "dearness_allowance",
    "daily_wage",
    "ssf_enrolled",
]
HEADERS_NE = [
    "पूरा नाम",
    "मोबाइल नं.",
    "पद",
    "सुरु मिति",
    "मिति प्रणाली (AD/BS)",
    "तलब प्रकार (monthly/daily)",
    "आधारभूत तलब",
    "महँगी भत्ता",
    "दैनिक ज्याला",
    "SSF (yes/no)",
]
TRUTHY = {"1", "true", "yes", "y", "हो", "x"}


def _rows_from_upload(upload) -> list[dict]:
    name = (upload.name or "").lower()
    raw = upload.read()
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(it, [])]
        rows = []
        for values in it:
            if values is None or all(v in (None, "") for v in values):
                rows.append(None)
                continue
            rows.append({header[i]: values[i] for i in range(min(len(header), len(values))) if header[i]})
        return rows
    if name.endswith(".csv") or not name:
        text = raw.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        reader.fieldnames = [str(f or "").strip().lower() for f in (reader.fieldnames or [])]
        return [
            r if any((v or "").strip() for v in r.values() if isinstance(v, str)) else None for r in reader
        ]
    raise APIError("validation_error", fields={"file": ["Upload a .csv or .xlsx file."]})


def _money(value):
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).replace(",", "").strip())
    except InvalidOperation as exc:
        raise ValueError(f"Invalid amount '{value}'") from exc


def _date(value, system):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    try:
        return parse_date_flexible(str(value or "").strip(), system or "AD")
    except BSDateError as exc:
        raise ValueError(f"Invalid start_date '{value}' ({system or 'AD'})") from exc


def import_employments(org, upload) -> dict:
    rows = _rows_from_upload(upload)
    created = updated = 0
    errors = []
    for idx, row in enumerate(rows, start=2):  # row 1 = header
        if row is None:
            continue
        try:
            with transaction.atomic():
                phone = normalize_phone(row.get("phone"))
                if not is_valid_phone(phone):
                    raise ValueError("Invalid phone number")
                wage_type = str(row.get("wage_type") or "monthly").strip().lower()
                if wage_type not in ("monthly", "daily"):
                    raise ValueError("wage_type must be monthly or daily")
                data = {
                    "designation": str(row.get("designation") or "").strip(),
                    "start_date": _date(row.get("start_date"), str(row.get("date_system") or "AD").strip()),
                    "wage_type": wage_type,
                    "basic_salary": _money(row.get("basic_salary")),
                    "dearness_allowance": _money(row.get("dearness_allowance")),
                    "daily_wage": _money(row.get("daily_wage")),
                    "ssf_enrolled": str(row.get("ssf_enrolled") or "").strip().lower() in TRUTHY,
                }
                worker, _ = get_or_create_worker(
                    {"phone": phone, "full_name": str(row.get("full_name") or "").strip()}
                )
                existing = Employment.objects.filter(organization=org, worker=worker, status="active").first()
                if existing:
                    validate_wage(
                        wage_type, data["basic_salary"], data["dearness_allowance"], data["daily_wage"]
                    )
                    for k, v in data.items():
                        if k in ("basic_salary", "dearness_allowance") and wage_type != "monthly":
                            v = None
                        if k == "daily_wage" and wage_type != "daily":
                            v = None
                        setattr(existing, k, v)
                    existing.save()
                    updated += 1
                else:
                    create_employment(org, worker, data)
                    created += 1
        except APIError as exc:
            msg = exc.message or str(exc.detail)
            if exc.fields:
                msg = "; ".join(f"{k}: {v[0]}" for k, v in exc.fields.items())
            errors.append({"row": idx, "message": str(msg)})
        except (ValueError, KeyError) as exc:
            errors.append({"row": idx, "message": str(exc)})
    return {"created": created, "updated": updated, "errors": errors}


def template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "employments"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2647C4")
    ws.append(
        ["Ram Bahadur Thapa", "9841000001", "Cook", "2081-04-01", "BS", "monthly", 12170, 7380, None, "yes"]
    )
    ws.append(["Sita Kumari Rai", "9812000002", "Helper", "2025-08-01", "AD", "daily", None, None, 800, "no"])
    for i, width in enumerate([24, 14, 16, 14, 12, 12, 14, 18, 12, 12], start=1):
        ws.column_dimensions[chr(64 + i)].width = width
    help_ws = wb.create_sheet("instructions / निर्देशन")
    help_ws.append(["column", "नेपाली", "notes"])
    notes = [
        "Worker's full name",
        "10-digit mobile (97/98…)",
        "Job title",
        "YYYY-MM-DD",
        "AD (default) or BS",
        "monthly or daily",
        "Monthly only",
        "Monthly only (basic+DA ≥ minimum wage)",
        "Daily only (≥ minimum daily wage)",
        "yes / no",
    ]
    for col, ne, note in zip(COLUMNS, HEADERS_NE, notes, strict=True):
        help_ws.append([col, ne, note])
    help_ws.column_dimensions["A"].width = 20
    help_ws.column_dimensions["B"].width = 28
    help_ws.column_dimensions["C"].width = 44
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
