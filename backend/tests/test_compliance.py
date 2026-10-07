"""Compliance: contracts, SSF filing, compliance status, contractors (CONTRACT-ECOSYSTEM §5)."""

import datetime as dt
import io
from decimal import Decimal as D
from unittest import mock

import pytest
from django.test import override_settings
from openpyxl import load_workbook

from apps.compliance import services
from apps.compliance.models import Contractor, EmploymentContract, SSFFiling
from apps.notifications.models import SMSMessage
from apps.payroll.models import PayrollRun, Payslip
from tests import factories as f
from tests.conftest import auth_client

pytestmark = pytest.mark.django_db

# Shrawan 2083 = 2026-07-17 .. 2026-08-16 (31 days)
SHRAWAN_START, SHRAWAN_END = dt.date(2026, 7, 17), dt.date(2026, 8, 16)


def make_run(org, status="approved", bs_month=4):
    run = PayrollRun.objects.create(
        organization=org,
        bs_year=2083,
        bs_month=bs_month,
        period_start=SHRAWAN_START,
        period_end=SHRAWAN_END,
        days_in_month=31,
        status=status,
    )
    return run


def make_slip(run, emp, basic="12170.00"):
    basic = D(basic)
    ssf_emp = (basic * D("0.11")).quantize(D("0.01")) if emp.ssf_enrolled else D("0")
    ssf_er = (basic * D("0.20")).quantize(D("0.01")) if emp.ssf_enrolled else D("0")
    return Payslip.objects.create(
        run=run,
        employment=emp,
        worker=emp.worker,
        wage_type=emp.wage_type,
        days_in_month=31,
        basic_earned=basic,
        da_earned=D("7380.00"),
        gross_pay=basic + D("7380.00"),
        ssf_employee=ssf_emp,
        ssf_employer=ssf_er,
        net_pay=basic + D("7380.00") - ssf_emp,
        payment_method=emp.worker.payment_method,
    )


# ---------------------------------------------------------------- SSF filing


def test_ssf_due_date_is_bs_month_end_plus_setting(org):
    run = make_run(org)
    with override_settings(SSF_DUE_DAYS=15):
        assert services.ssf_due_date(run) == SHRAWAN_END + dt.timedelta(days=15)


def test_ssf_rows_and_missing_sheet(owner_client, org):
    run = make_run(org)
    with_id = f.EmploymentFactory(
        organization=org,
        ssf_enrolled=True,
        worker=f.WorkerFactory(full_name="Asha", ssf_id="SSF123", citizenship_no="27-01-123"),
    )
    no_id = f.EmploymentFactory(
        organization=org, ssf_enrolled=True, worker=f.WorkerFactory(full_name="Bina", ssf_id=None)
    )
    not_enrolled = f.EmploymentFactory(organization=org, ssf_enrolled=False)
    for e in (with_id, no_id, not_enrolled):
        make_slip(run, e)

    rows, missing = services.ssf_rows(run)
    assert [r["ssf_id"] for r in rows] == ["SSF123"]
    r = rows[0]
    assert r["basic_earned"] == "12170.00" and r["employee_11"] == "1338.70"
    assert r["employer_20"] == "2434.00" and r["total_31"] == "3772.70"
    assert r["bs_year"] == 2083 and r["bs_month"] == 4 and r["citizenship_no"] == "27-01-123"
    assert [m["worker_name"] for m in missing] == ["Bina"]

    res = owner_client.post(f"/api/v1/org/payroll-runs/{run.id}/ssf-filing")
    assert res.status_code == 201, res.json()
    body = res.json()
    assert body["status"] == "prepared"
    assert body["due_date"] == (SHRAWAN_END + dt.timedelta(days=15)).isoformat()
    assert body["due_date_bs"]
    assert body["summary"]["rows"] == 1 and body["summary"]["missing"] == 1
    assert body["summary"]["total_31"] == "3772.70"

    xlsx = owner_client.get(f"/api/v1/org/ssf-filings/{body['id']}/file.xlsx")
    assert xlsx.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["SSF contributions", "Missing SSF ID"]
    ws = wb["SSF contributions"]
    assert [c.value for c in ws[3]] == services.SSF_COLUMNS
    assert ws.cell(row=4, column=1).value == "SSF123"
    assert wb["Missing SSF ID"].cell(row=4, column=1).value == "Bina"

    csv_res = owner_client.get(f"/api/v1/org/ssf-filings/{body['id']}/file.csv")
    lines = csv_res.content.decode().strip().splitlines()
    assert lines[0] == ",".join(services.SSF_COLUMNS)
    assert lines[1].startswith("SSF123,Asha,27-01-123,12170.00,1338.70,2434.00,3772.70,2083,4")


def test_ssf_filing_requires_approved_run(owner_client, org):
    run = make_run(org, status="draft")
    res = owner_client.post(f"/api/v1/org/payroll-runs/{run.id}/ssf-filing")
    assert res.status_code == 400 and res.json()["error"]["code"] == "invalid_state"


def test_ssf_filing_status_flow(owner_client, org):
    run = make_run(org, status="paid")
    filing, _ = services.prepare_ssf_filing(run)
    url = f"/api/v1/org/ssf-filings/{filing.id}"
    res = owner_client.patch(url, {"status": "submitted"}, format="json")
    assert res.status_code == 200 and res.json()["submitted_on"]
    # Re-preparing a submitted filing is not allowed.
    assert owner_client.post(f"/api/v1/org/payroll-runs/{run.id}/ssf-filing").status_code == 400
    res = owner_client.patch(url, {"status": "acknowledged"}, format="json")
    assert res.status_code == 400 and "receipt_no" in res.json()["error"]["fields"]
    res = owner_client.patch(url, {"status": "acknowledged", "receipt_no": "R-77"}, format="json")
    assert res.status_code == 200 and res.json()["receipt_no"] == "R-77"
    res = owner_client.patch(url, {"status": "prepared"}, format="json")
    assert res.json()["error"]["code"] == "invalid_state"
    lst = owner_client.get("/api/v1/org/ssf-filings").json()
    assert lst["count"] == 1 and lst["results"][0]["run"]["bs_month"] == 4


def test_ssf_reminders_sms_5_and_1_days_before(org):
    from apps.compliance.tasks import ssf_reminders

    owner = f.MembershipFactory(organization=org, role="owner").user
    f.MembershipFactory(organization=org, role="supervisor")
    run = make_run(org)
    make_slip(run, f.EmploymentFactory(organization=org, ssf_enrolled=True))
    due = services.ssf_due_date(run)
    for days, expected in ((5, 1), (3, 0), (1, 1)):
        SMSMessage.objects.all().delete()
        with mock.patch("apps.core.bs.today_ad", return_value=due - dt.timedelta(days=days)):
            with pytest.MonkeyPatch.context() as mp:
                mp.setattr("django.db.transaction.on_commit", lambda fn, *a, **k: fn())
                assert ssf_reminders() == expected
        assert SMSMessage.objects.filter(phone=owner.phone, kind="ssf_reminder").count() == expected
    # Submitted filings stop reminders.
    filing, _ = services.prepare_ssf_filing(run)
    services.update_filing(filing, {"status": "submitted"})
    with mock.patch("apps.core.bs.today_ad", return_value=due - dt.timedelta(days=1)):
        assert ssf_reminders() == 0


# ---------------------------------------------------------------- contracts


def test_contract_versioning_and_pdf(owner_client, org, site):
    emp = f.EmploymentFactory(organization=org, site=site, designation="Cook")
    url = f"/api/v1/org/employments/{emp.id}/contract"
    assert owner_client.get(url).json()["latest"] is None

    res = owner_client.post(url, {}, format="json")
    assert res.status_code == 201
    assert res.json()["latest"]["version"] == 1 and res.json()["latest"]["is_accepted"] is False
    # Unchanged terms → same version.
    res = owner_client.post(url, {}, format="json")
    assert res.status_code == 200 and len(res.json()["versions"]) == 1

    emp.basic_salary = D("15000.00")
    emp.save()
    assert owner_client.get(url).json()["needs_new_version"] is True
    res = owner_client.post(url, {}, format="json")
    assert res.status_code == 201
    data = res.json()
    assert data["latest"]["version"] == 2 and data["latest"]["terms"]["basic_salary"] == "15000.00"
    assert [v["version"] for v in data["versions"]] == [2, 1]
    # Forced regeneration also creates a version.
    assert owner_client.post(url, {"force": True}, format="json").json()["latest"]["version"] == 3

    pdf = owner_client.get(f"{url}.pdf")
    assert pdf.status_code == 200 and pdf["Content-Type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF")
    old = owner_client.get(f"{url}.pdf?version=1")
    assert old.status_code == 200 and "v1" in old["Content-Disposition"]


def test_contract_permissions(org, site, make_member):
    emp = f.EmploymentFactory(organization=org, site=site)
    url = f"/api/v1/org/employments/{emp.id}/contract"
    _, sup = make_member("supervisor", sites=[site])
    assert sup.get(url).status_code == 403
    _, hr = make_member("hr")
    assert hr.post(url, {}, format="json").status_code == 201
    other = f.OrganizationFactory()
    outsider = f.MembershipFactory(organization=other, role="owner").user
    assert auth_client(outsider, other).get(url).status_code == 404


def test_worker_accepts_latest_contract(org):
    user = f.UserFactory()
    worker = f.WorkerFactory(user=user, phone=user.phone)
    emp = f.EmploymentFactory(organization=org, worker=worker)
    client = auth_client(user)
    base = f"/api/v1/me/employments/{emp.id}/contract"
    assert client.get(base).status_code == 404
    assert client.post(f"{base}/accept").status_code == 404

    services.generate_contract(emp)
    res = client.get(base)
    assert res.status_code == 200 and res.json()["version"] == 1
    pdf = client.get(f"{base}.pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    res = client.post(f"{base}/accept")
    assert res.status_code == 200 and res.json()["is_accepted"] and res.json()["accepted_channel"] == "app"

    # A wage change creates an unaccepted v2.
    emp.other_allowance = D("2000.00")
    emp.save()
    services.generate_contract(emp)
    assert client.get(base).json()["is_accepted"] is False
    # Another worker cannot see it.
    stranger = f.UserFactory()
    f.WorkerFactory(user=stranger, phone=stranger.phone)
    assert auth_client(stranger).get(base).status_code == 404
    assert auth_client(stranger).post(f"{base}/accept").status_code == 404
    assert EmploymentContract.objects.get(employment=emp, version=1).accepted_at is not None


def test_org_records_paper_acceptance(owner_client, org):
    emp = f.EmploymentFactory(organization=org)
    services.generate_contract(emp)
    res = owner_client.post(
        f"/api/v1/org/employments/{emp.id}/contract/accept", {"channel": "paper"}, format="json"
    )
    assert res.status_code == 200 and res.json()["latest"]["accepted_channel"] == "paper"


# ---------------------------------------------------------------- compliance status


def items_by_key(body):
    return {i["key"]: i for i in body["items"]}


def test_compliance_items(owner_client, org):
    org.ssf_employer_no = None
    org.save()
    good = f.EmploymentFactory(organization=org, ssf_enrolled=True, worker=f.WorkerFactory(ssf_id="S1"))
    f.EmploymentFactory(organization=org, ssf_enrolled=True, worker=f.WorkerFactory(ssf_id=None))
    low = f.EmploymentFactory(organization=org)
    type(low).objects.filter(pk=low.pk).update(basic_salary=D("9000.00"))
    services.generate_contract(good)

    res = owner_client.get("/api/v1/org/compliance")
    assert res.status_code == 200
    body = res.json()
    items = items_by_key(body)
    assert set(items) == set(services.WEIGHTS)
    assert items["ssf_employer_registered"]["status"] == "action"
    assert items["workers_missing_ssf_id"] == {
        **items["workers_missing_ssf_id"],
        "status": "warning",
        "count": 1,
    }
    assert items["below_minimum_wage"]["status"] == "action" and items["below_minimum_wage"]["count"] == 1
    assert items["contracts_missing"]["count"] == 2
    assert items["contracts_unaccepted"]["count"] == 1
    assert items["payroll_not_run_last_month"]["status"] == "action"
    assert items["attendance_unmarked_7d"]["count"] > 0
    assert items["credentials_expiring"]["status"] == "ok"
    assert all(i["link"].startswith("/app/") and i["message"] for i in body["items"])
    assert 0 <= body["score"] < 60


def test_compliance_clean_org_scores_high(owner_client, org):
    org.ssf_employer_no = "SSF-ER-1"
    org.save()
    res = owner_client.get("/api/v1/org/compliance")
    body = res.json()
    assert body["score"] == 100, body
    assert all(i["status"] == "ok" for i in body["items"])


def test_ssf_overdue_item(owner_client, org):
    org.ssf_employer_no = "SSF-ER-1"
    org.save()
    run = make_run(org, status="paid")
    make_slip(
        run, f.EmploymentFactory(organization=org, ssf_enrolled=True, worker=f.WorkerFactory(ssf_id="S9"))
    )
    items = items_by_key(owner_client.get("/api/v1/org/compliance").json())
    assert items["ssf_filings_overdue"]["status"] == "action" and items["ssf_filings_overdue"]["count"] == 1
    filing, _ = services.prepare_ssf_filing(run)
    services.update_filing(filing, {"status": "submitted"})
    items = items_by_key(owner_client.get("/api/v1/org/compliance").json())
    assert items["ssf_filings_overdue"]["status"] == "ok"


def test_compliance_is_manager_only(org, make_member):
    _, sup = make_member("supervisor")
    assert sup.get("/api/v1/org/compliance").status_code == 403


# ---------------------------------------------------------------- contractors


def test_contractor_crud_assign_and_summary(owner_client, org):
    res = owner_client.post(
        "/api/v1/org/contractors",
        {"name": "Ram Thekedar", "phone": "9841000000", "pan_no": "123456789", "trades": ["mason", "helper"]},
        format="json",
    )
    assert res.status_code == 201, res.json()
    cid = res.json()["id"]
    assert {t["code"] for t in res.json()["trades"]} == {"mason", "helper"}

    e1 = f.EmploymentFactory(organization=org, start_date=dt.date(2026, 1, 1))
    e2 = f.EmploymentFactory(
        organization=org,
        start_date=dt.date(2026, 1, 1),
        wage_type="daily",
        basic_salary=None,
        dearness_allowance=None,
        daily_wage=D("1000"),
    )
    f.EmploymentFactory(organization=org)
    res = owner_client.post(
        f"/api/v1/org/contractors/{cid}/assign", {"employment_ids": [str(e1.id), str(e2.id)]}, format="json"
    )
    assert res.status_code == 200 and res.json()["active_workers"] == 2

    lst = owner_client.get(f"/api/v1/org/employments?contractor={cid}").json()
    assert {r["id"] for r in lst["results"]} == {str(e1.id), str(e2.id)}

    for d in (dt.date(2026, 7, 20), dt.date(2026, 7, 21)):
        f.AttendanceFactory(employment=e2, date=d)
    f.AttendanceFactory(employment=e2, date=dt.date(2026, 7, 22), status="half_day")
    run = make_run(org)
    make_slip(run, e1)
    s = owner_client.get(f"/api/v1/org/contractors/{cid}/summary?bs_year=2083&bs_month=4").json()
    assert s["headcount"] == 2 and s["payroll_run"]["id"] == str(run.id)
    assert s["wages"]["gross"] == "19550.00"
    by_emp = {w["employment_id"]: w for w in s["workers"]}
    assert by_emp[str(e2.id)]["days"] == "2.5" and by_emp[str(e2.id)]["gross_pay"] is None

    other_org_emp = f.EmploymentFactory()
    res = owner_client.post(
        f"/api/v1/org/contractors/{cid}/assign", {"employment_ids": [str(other_org_emp.id)]}, format="json"
    )
    assert res.status_code == 400

    res = owner_client.patch(f"/api/v1/org/contractors/{cid}", {"is_active": False}, format="json")
    assert res.status_code == 200 and res.json()["is_active"] is False
    assert owner_client.delete(f"/api/v1/org/contractors/{cid}").status_code == 204
    assert not Contractor.objects.filter(pk=cid).exists()


def test_contractors_scoped_to_org(owner_client, org):
    Contractor.objects.create(organization=f.OrganizationFactory(), name="Other", phone="9800000000")
    assert owner_client.get("/api/v1/org/contractors").json()["count"] == 0


def test_ssf_filing_model_one_per_run(org):
    run = make_run(org)
    services.prepare_ssf_filing(run)
    filing, created = services.prepare_ssf_filing(run)
    assert created is False and SSFFiling.objects.filter(run=run).count() == 1


def test_messages_in_nepali(owner_client, org):
    org.ssf_employer_no = None
    org.save()
    items = items_by_key(owner_client.get("/api/v1/org/compliance", HTTP_ACCEPT_LANGUAGE="ne").json())
    assert "सामाजिक सुरक्षा कोष" in items["ssf_employer_registered"]["message"]
    run = make_run(org)
    res = owner_client.post(f"/api/v1/org/payroll-runs/{run.id}/ssf-filing", HTTP_ACCEPT_LANGUAGE="ne")
    assert res.json()["status_label"] == "तयार भयो"
